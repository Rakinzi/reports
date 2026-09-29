"""Offline checks that fresh evidence replaces the edited report templates."""

import json
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation

from reports.generator_2026 import (
    _SLIDE1_KPI_CARD_PICTURE,
    _SLIDE3_SNAPSHOT_CARD_PICTURE,
    _SLIDE4_COUNTRIES_TABLE_PICTURE,
    _SLIDE5_PAGES_TABLE_PICTURE,
    _SLIDE6_SEARCH_CONSOLE_PICTURE,
    _SLIDE_TRAFFIC_ACQ_PICTURE,
    _TOP_QUERIES_PICTURE,
    _SECURITY_HEADERS_PICTURE,
    TEMPLATES_2026,
    TEMPLATES_DIR,
    _apply_supplemental_evidence,
    _build_slide_security,
    _find_slide_index,
    _fill_paragraph_slots,
    _remove_duplicate_security_slides,
    _required_captures,
    _seo_picture_slots,
    _seo_result_narrative,
    _sitecheck_picture_slot,
    _validate_core_metrics,
    _verify_embedded_captures,
)
from reports.template_runner import _fill_template_sections
import reports.generator_2026 as generator


def _image(path: Path, color: tuple[int, int, int]) -> Path:
    Image.new("RGB", (80, 50), color).save(path)
    return path


@pytest.mark.parametrize("report_name", sorted(TEMPLATES_2026))
def test_recent_templates_replace_all_seo_and_security_evidence(report_name, tmp_path):
    prs = Presentation(str(TEMPLATES_DIR / TEMPLATES_2026[report_name]))
    for title, names in (
        (None, _SLIDE1_KPI_CARD_PICTURE),
        ("Site Overview", _SLIDE3_SNAPSHOT_CARD_PICTURE),
        ("Geographic Performance", _SLIDE4_COUNTRIES_TABLE_PICTURE),
        ("Page Performance", _SLIDE5_PAGES_TABLE_PICTURE),
        ("Traffic Acquisition", _SLIDE_TRAFFIC_ACQ_PICTURE),
        ("Search Performance", _SLIDE6_SEARCH_CONSOLE_PICTURE),
        ("Top Queries", _TOP_QUERIES_PICTURE),
        ("Security", _SECURITY_HEADERS_PICTURE),
    ):
        slide_index = 0 if title is None else _find_slide_index(prs, title)
        if slide_index is not None:
            picture_names = {s.name for s in prs.slides[slide_index].shapes if s.shape_type == 13}
            assert names[report_name] in picture_names, (report_name, title, picture_names)
    has_uptime_before = any(
        any(s.has_text_frame and s.text_frame.text.strip().casefold().startswith("uptime")
            for s in slide.shapes)
        for slide in prs.slides
    )
    _remove_duplicate_security_slides(prs)
    assert any(
        any(s.has_text_frame and s.text_frame.text.strip().casefold().startswith("uptime")
            for s in slide.shapes)
        for slide in prs.slides
    ) == has_uptime_before

    screenshots = {}
    security_index = _find_slide_index(prs, "Security")
    security_narrative_before = None
    if security_index is not None:
        security_narrative_before = next(
            (shape.text for shape in prs.slides[security_index].shapes
             if shape.name == "object 5" and shape.has_text_frame),
            None,
        )
        screenshots["security_headers_screenshot"] = _image(tmp_path / "headers.png", (22, 71, 136))
        _build_slide_security(prs.slides[security_index], screenshots, report_name)
        _verify_embedded_captures(
            prs, screenshots, {security_index: "security_headers_screenshot"}, report_name
        )
        if security_narrative_before is not None:
            security_narrative_after = next(
                shape.text for shape in prs.slides[security_index].shapes
                if shape.name == "object 5" and shape.has_text_frame
            )
            assert security_narrative_after == security_narrative_before

    slots = _seo_picture_slots(prs, report_name)
    for index, (_, _, query, key) in enumerate(slots):
        screenshots[key] = _image(tmp_path / f"{key}.png", (index + 1, 40, 90))
        screenshots[key].with_suffix(".json").write_text(json.dumps({
            "query": query,
            "results": [{"title": "Example result", "domain": "example.com"}],
        }), encoding="utf-8")
    sitecheck_slot = _sitecheck_picture_slot(prs, report_name)
    if sitecheck_slot:
        screenshots["sitecheck_screenshot"] = _image(tmp_path / "sitecheck.png", (35, 134, 118))

    _apply_supplemental_evidence(prs, report_name, screenshots)
    for slide_index, _, _, key in slots:
        _verify_embedded_captures(prs, screenshots, {slide_index: key}, report_name)
    if sitecheck_slot:
        _verify_embedded_captures(
            prs, screenshots, {sitecheck_slot[0]: "sitecheck_screenshot"}, report_name
        )
    if security_index is not None and security_narrative_before is not None:
        security_narrative_after = next(
            shape.text for shape in prs.slides[security_index].shapes
            if shape.name == "object 5" and shape.has_text_frame
        )
        assert security_narrative_after == security_narrative_before


def test_ga4_new_users_can_exceed_active_users_without_blocking_report():
    _validate_core_metrics({
        "Active users": "8,137", "New users": "8,500",
        "Average engagement time per active user": "49s",
    }, "zimplats")
    with pytest.raises(RuntimeError, match="Active users is missing"):
        _validate_core_metrics({"New users": "8,500"}, "zimplats")


def test_removed_template_text_is_not_reintroduced():
    prs = Presentation(str(TEMPLATES_DIR / TEMPLATES_2026["bancabc"]))
    slide = prs.slides[_find_slide_index(prs, "Site Overview")]
    shape = next(s for s in slide.shapes if s.name == "object 2")
    for paragraph in shape.text_frame.paragraphs:
        paragraph.text = ""
    assert not _fill_paragraph_slots(shape, ["replacement"])
    assert all(paragraph.text == "" for paragraph in shape.text_frame.paragraphs)


def test_ecocash_keeps_supported_edited_search_insight():
    prs = Presentation(str(TEMPLATES_DIR / TEMPLATES_2026["ecocash"]))
    slide = prs.slides[_find_slide_index(prs, "SEO")]
    existing = next(shape.text_frame.text for shape in slide.shapes if shape.name == "object 5")
    assert _seo_result_narrative(
        "ecocash", "EcoCash",
        [{"title": "EcoCash Zimbabwe", "domain": "ecocash.co.zw"}],
        existing,
    ) is None
    assert _seo_result_narrative(
        "ecocash", "EcoCash",
        [{"title": "Third-party page", "domain": "example.com"}],
        existing,
    ) is not None


def test_missing_current_capture_blocks_builtin_report(tmp_path):
    prs = Presentation(str(TEMPLATES_DIR / TEMPLATES_2026["bancabc"]))
    screenshot = _image(tmp_path / "snapshot.png", (1, 2, 3))
    with pytest.raises(RuntimeError, match="countries_table"):
        _required_captures(
            prs, {"snapshot_card": screenshot},
            [{"country": "Zimbabwe"}], [{"path": "/"}],
            [{"source_medium": "google / organic"}],
            {"impressions": "1", "clicks": "1", "ctr": "100%", "avg_position": "1",
             "top_queries": [{"query": "brand", "clicks": 1}]},
            "bancabc",
        )


def test_uploaded_template_does_not_keep_an_old_mapped_image(tmp_path):
    source = TEMPLATES_DIR / TEMPLATES_2026["bancabc"]
    output = tmp_path / "report.pptx"
    output.write_bytes(source.read_bytes())
    mapping = [{
        "slide_index": 4,
        "shape_name": "Picture 9",
        "shape_type": "image",
        "field_type": "screenshot_countries_table",
    }]
    property_section = [{"section_key": None, "start_slide": 0, "end_slide": 99}]
    with pytest.raises(RuntimeError, match="Current image.*missing"):
        _fill_template_sections(output, mapping, property_section, [], {None: {}}, {None: {}})


@pytest.mark.parametrize("report_name", sorted(TEMPLATES_2026))
def test_every_builtin_deck_builds_with_fresh_distinct_captures(report_name, tmp_path, monkeypatch):
    template = Presentation(str(TEMPLATES_DIR / TEMPLATES_2026[report_name]))
    keys = {
        "snapshot_card", "countries_table", "pages_table", "traffic_acquisition_table",
        "search_screenshot", "top_queries_screenshot", "security_headers_screenshot",
        "sitecheck_screenshot",
    }
    keys.update(key for _, _, _, key in _seo_picture_slots(template, report_name))
    captures = {
        key: _image(tmp_path / f"{key}.png", (index + 1, 80, 130))
        for index, key in enumerate(sorted(keys))
    }
    for _, _, query, key in _seo_picture_slots(template, report_name):
        captures[key].with_suffix(".json").write_text(json.dumps({
            "query": query,
            "results": [{"title": "Example result", "domain": "example.com"}],
        }), encoding="utf-8")
    pages = [{
        "title": "Homepage", "path": "/", "views": 1000,
        "views_pct": "70%", "active_users": 500,
        "views_per_user": "2.0", "avg_engagement_time": "49s",
    }]
    countries = [{
        "country": "Zimbabwe", "users": 500, "new_users": 400,
        "engaged_sessions": 300, "engagement_rate": "50%",
        "engaged_sessions_per_user": "0.6",
    }]
    traffic = [{
        "source_medium": "google / organic", "sessions": 600,
        "sessions_pct": "60%", "engaged_sessions": 300,
        "engaged_sessions_pct": "60%", "engagement_rate": "50%",
        "avg_engagement_time": "30s", "events_per_session": "3.0",
    }]
    search = {
        "impressions": "10K", "clicks": "500", "ctr": "5%",
        "avg_position": "3", "top_queries": [{"query": "brand", "clicks": 100}],
    }
    result = (
        captures,
        {"Active users": "500", "New users": "400", "Average engagement time per active user": "49s"},
        {"countries": {"Zimbabwe": 500}, "channels": {"Organic Search": 600}},
        {"/": 1000}, pages, 1000, countries, search, traffic,
        {"sessions": 1000, "engaged_sessions": 500, "engagement_rate": "50%"},
    )
    monkeypatch.setattr(generator, "capture_2026", lambda *args, **kwargs: result)
    monkeypatch.setattr(generator, "_gemini_paras_batch", lambda values: values)
    monkeypatch.setattr(generator, "OUTPUT_DIR", tmp_path)
    output = generator.generate_report_2026(
        report_name, "1 September 2026 - 28 September 2026", "29 September 2026",
        "Sep 1, 2026", "Sep 28, 2026",
    )
    assert output.is_file()
    built = Presentation(str(output))
    _remove_duplicate_security_slides(template)
    for source_slide, built_slide in zip(template.slides, built.slides):
        built_text = {shape.name: shape for shape in built_slide.shapes if shape.has_text_frame}
        for source_shape in source_slide.shapes:
            if not source_shape.has_text_frame or source_shape.name not in built_text:
                continue
            source_paras = source_shape.text_frame.paragraphs
            output_paras = built_text[source_shape.name].text_frame.paragraphs
            for index, paragraph in enumerate(source_paras):
                if not paragraph.text.strip() and index < len(output_paras):
                    assert not output_paras[index].text.strip(), (
                        report_name, source_shape.name, index, output_paras[index].text
                    )
    old_site = _find_slide_index(template, "Site Overview")
    new_site = _find_slide_index(built, "Site Overview")
    if old_site is not None and new_site is not None:
        old_subtitle = next((s for s in template.slides[old_site].shapes if s.name == "object 2"), None)
        new_subtitle = next((s for s in built.slides[new_site].shapes if s.name == "object 2"), None)
        if old_subtitle and new_subtitle and len(old_subtitle.text_frame.paragraphs) > 1:
            old_text = old_subtitle.text_frame.paragraphs[1].text.strip()
            if old_text:
                assert new_subtitle.text_frame.paragraphs[1].text.strip() != old_text
