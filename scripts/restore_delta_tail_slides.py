"""Restore Delta's Uptime/Downtime and empty Recommendations slides."""

from copy import deepcopy
from io import BytesIO
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


ROOT = Path(__file__).resolve().parents[1]
TARGET = ROOT / "src/reports/report-templates/june-2026/delta-25-June-2026.pptx"
SOURCE = Path("/Users/rakinzisilver/Downloads/delta-27-July-2026.pptx")


def _texts(slide) -> list[str]:
    return [
        shape.text.strip()
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False) and shape.text.strip()
    ]


def _clone_slide(destination, source_slide):
    layout = min(destination.slide_layouts, key=lambda item: len(item.placeholders))
    slide = destination.slides.add_slide(layout)
    for shape in list(slide.shapes):
        slide.shapes._spTree.remove(shape._element)
    for shape in source_slide.shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            picture = slide.shapes.add_picture(
                BytesIO(shape.image.blob), shape.left, shape.top, shape.width, shape.height
            )
            picture.name = shape.name
        else:
            slide.shapes._spTree.insert_element_before(deepcopy(shape._element), "p:extLst")
    return slide


def main() -> None:
    target = Presentation(TARGET)
    source = Presentation(SOURCE)

    uptime = next(
        slide for slide in source.slides
        if any("uptime and system reliability" in text.casefold() for text in _texts(slide))
    )
    recommendations = next(
        slide for slide in source.slides
        if any(text.casefold() == "recommendations" for text in _texts(slide))
    )

    if not any("uptime and system reliability" in text.casefold() for slide in target.slides for text in _texts(slide)):
        _clone_slide(target, uptime)
    if not any(text.casefold() == "recommendations" for slide in target.slides for text in _texts(slide)):
        rec_slide = _clone_slide(target, recommendations)
        for shape in rec_slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            if shape.text.strip().casefold() != "recommendations":
                shape.text = ""

    target.save(TARGET)
    print(f"Restored Delta tail slides; slide count={len(target.slides)}")


if __name__ == "__main__":
    main()
