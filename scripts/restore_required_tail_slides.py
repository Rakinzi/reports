"""Ensure every active June template ends with Uptime and Recommendations."""

from copy import deepcopy
from io import BytesIO
from pathlib import Path
from tempfile import NamedTemporaryFile
from zipfile import ZIP_DEFLATED, ZipFile

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE_DIR = ROOT / "src/reports/report-templates/june-2026"
REFERENCE = Path("/Users/rakinzisilver/Downloads/infraco-24-July-2026.pptx")


def _texts(slide) -> list[str]:
    return [
        shape.text.strip()
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False) and shape.text.strip()
    ]


def _is_uptime(slide) -> bool:
    return any(
        "uptime" in text.casefold() and "reliability" in text.casefold()
        for text in _texts(slide)
    )


def _is_recommendations(slide) -> bool:
    return any(text.casefold() == "recommendations" for text in _texts(slide))


def _clone_slide(destination, source_slide):
    layout = min(destination.slide_layouts, key=lambda item: len(item.placeholders))
    slide = destination.slides.add_slide(layout)
    for shape in list(slide.shapes):
        slide.shapes._spTree.remove(shape._element)
    for shape in source_slide.shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            picture = slide.shapes.add_picture(
                BytesIO(shape.image.blob), shape.left, shape.top, shape.width, shape.height
            )
            picture.name = shape.name
        else:
            slide.shapes._spTree.insert_element_before(deepcopy(shape._element), "p:extLst")
    return slide


def _clear_recommendations(slide) -> None:
    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        if shape.text.strip().casefold() != "recommendations":
            shape.text = ""


def _dedupe_package(path: Path) -> None:
    with ZipFile(path, "r") as source:
        entries = {info.filename: source.read(info) for info in source.infolist()}
    with NamedTemporaryFile(suffix=".pptx", dir=path.parent, delete=False) as handle:
        temporary = Path(handle.name)
    try:
        with ZipFile(temporary, "w", compression=ZIP_DEFLATED) as destination:
            for name, payload in entries.items():
                destination.writestr(name, payload)
        temporary.replace(path)
    finally:
        if temporary.exists():
            temporary.unlink()


def main() -> None:
    reference = Presentation(REFERENCE)
    uptime_source = next(slide for slide in reference.slides if _is_uptime(slide))
    recommendations_source = next(
        slide for slide in reference.slides if _is_recommendations(slide)
    )

    for path in sorted(TEMPLATE_DIR.glob("*.pptx")):
        presentation = Presentation(path)
        uptime = next((slide for slide in presentation.slides if _is_uptime(slide)), None)
        if uptime is None:
            uptime = _clone_slide(presentation, uptime_source)

        recommendations = next(
            (slide for slide in presentation.slides if _is_recommendations(slide)), None
        )
        if recommendations is None:
            recommendations = _clone_slide(presentation, recommendations_source)
        _clear_recommendations(recommendations)

        # Place the two required pages at the tail in the required order.
        for slide in (uptime, recommendations):
            slide_id = next(
                item for item in presentation.slides._sldIdLst
                if presentation.part.related_part(item.rId) is slide.part
            )
            presentation.slides._sldIdLst.remove(slide_id)
            presentation.slides._sldIdLst.append(slide_id)

        presentation.save(path)
        _dedupe_package(path)
        print(f"updated {path.name}: Uptime then Recommendations")


if __name__ == "__main__":
    main()
