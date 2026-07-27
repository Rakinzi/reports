"""Use the approved Infraco Overview as slide 2 in every report template."""

from copy import deepcopy
from pathlib import Path
from tempfile import NamedTemporaryFile
from zipfile import ZIP_DEFLATED, ZipFile

from pptx import Presentation


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE_DIR = ROOT / "src/reports/report-templates/june-2026"
REFERENCE_TEMPLATE = Path("/Users/rakinzisilver/Downloads/infraco-24-July-2026.pptx")


def _dedupe_package(path: Path) -> None:
    """Rewrite a PPTX with one ZIP entry per part, keeping the newest entry."""
    with ZipFile(path, "r") as source:
        entries = {info.filename: source.read(info) for info in source.infolist()}
    with NamedTemporaryFile(suffix=".pptx", dir=path.parent, delete=False) as handle:
        temporary = Path(handle.name)
    try:
        with ZipFile(temporary, "w", compression=ZIP_DEFLATED) as destination:
            for name, payload in entries.items():
                destination.writestr(name, payload)
        temporary.replace(path)
    finally:
        if temporary.exists():
            temporary.unlink()


def _slide_index(prs: Presentation, title: str) -> int | None:
    wanted = title.casefold()
    for index, slide in enumerate(prs.slides):
        if any(
            getattr(shape, "has_text_frame", False)
            and shape.text_frame.text.strip().casefold().startswith(wanted)
            for shape in slide.shapes
        ):
            return index
    return None


def _slide_indices(prs: Presentation, title: str) -> list[int]:
    """Return every slide whose title exactly matches ``title``."""
    wanted = title.casefold()
    matches: list[int] = []
    for index, slide in enumerate(prs.slides):
        if any(
            getattr(shape, "has_text_frame", False)
            and shape.text_frame.text.strip().casefold() == wanted
            for shape in slide.shapes
        ):
            matches.append(index)
    return matches


def _move_slide(prs: Presentation, source_index: int, destination_index: int) -> None:
    slide_id = prs.slides._sldIdLst[source_index]
    prs.slides._sldIdLst.remove(slide_id)
    prs.slides._sldIdLst.insert(destination_index, slide_id)


def _remove_slide(prs: Presentation, slide_index: int) -> None:
    slide_id = prs.slides._sldIdLst[slide_index]
    prs.part.drop_rel(slide_id.rId)
    prs.slides._sldIdLst.remove(slide_id)


def _replace_overview_shapes(destination_slide, reference_slide) -> None:
    """Replace shapes in an existing slide, avoiding slide-part name collisions."""
    for shape in list(destination_slide.shapes):
        destination_slide.shapes._spTree.remove(shape._element)
    for shape in reference_slide.shapes:
        destination_slide.shapes._spTree.insert_element_before(
            deepcopy(shape._element), "p:extLst"
        )

    narrative = next(
        (shape for shape in destination_slide.shapes if shape.name == "object 8"), None
    )
    if narrative is not None:
        for paragraph in narrative.text_frame.paragraphs:
            paragraph.text = ""

def main() -> None:
    reference = Presentation(REFERENCE_TEMPLATE)
    reference_index = _slide_index(reference, "Overview")
    if reference_index is None:
        raise RuntimeError("Reference template has no Overview page")
    reference_slide = reference.slides[reference_index]

    for template_path in sorted(TEMPLATE_DIR.glob("*.pptx")):
        presentation = Presentation(template_path)
        overview_indices = _slide_indices(presentation, "Overview")
        if not overview_indices:
            raise RuntimeError(f"{template_path.name} has no Overview page to standardize")

        # Keep one existing slide part and remove every duplicate. Reusing the
        # part avoids python-pptx assigning a colliding slide filename.
        keep_index = overview_indices[0]
        for overview_index in reversed(overview_indices[1:]):
            _remove_slide(presentation, overview_index)
        keep_index = _slide_indices(presentation, "Overview")[0]
        _replace_overview_shapes(presentation.slides[keep_index], reference_slide)
        _move_slide(presentation, keep_index, 1)
        presentation.save(template_path)
        _dedupe_package(template_path)
        print(f"replaced {template_path.name}: Infraco Overview is slide 2")


if __name__ == "__main__":
    main()
