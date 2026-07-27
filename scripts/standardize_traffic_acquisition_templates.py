"""Insert the standard Traffic Acquisition slide into June 2026 templates.

The reference layout comes from the approved Infraco July report. Existing
templates are left unchanged when they already contain a Traffic Acquisition
page. The inserted page is positioned immediately after Page Performance.
"""

from copy import deepcopy
from io import BytesIO
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE_DIR = ROOT / "src/reports/report-templates/june-2026"
REFERENCE_DECK = Path("/Users/rakinzisilver/Downloads/infraco-24-July-2026.pptx")


def _slide_index(prs: Presentation, title: str) -> int | None:
    wanted = title.casefold()
    for index, slide in enumerate(prs.slides):
        if any(
            getattr(shape, "has_text_frame", False)
            and shape.text_frame.text.strip().casefold().startswith(wanted)
            for shape in slide.shapes
        ):
            return index
    return None


def _clear_reference_data(slide) -> None:
    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        if shape.name == "object 3":
            for paragraph in shape.text_frame.paragraphs:
                paragraph.text = ""
        elif shape.name == "object 7":
            for index, paragraph in enumerate(shape.text_frame.paragraphs):
                paragraph.text = "Overall Insight:" if index == 0 else ""


def _insert_reference_slide(destination: Presentation, reference_slide) -> None:
    layout = min(destination.slide_layouts, key=lambda item: len(item.placeholders))
    new_slide = destination.slides.add_slide(layout)

    for shape in list(new_slide.shapes):
        new_slide.shapes._spTree.remove(shape._element)

    for shape in reference_slide.shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            picture = new_slide.shapes.add_picture(
                BytesIO(shape.image.blob),
                shape.left,
                shape.top,
                shape.width,
                shape.height,
            )
            picture.name = "Picture 32"
        else:
            new_slide.shapes._spTree.insert_element_before(
                deepcopy(shape._element), "p:extLst"
            )

    _clear_reference_data(new_slide)

    page_performance_index = _slide_index(destination, "Page Performance")
    if page_performance_index is None:
        raise RuntimeError("Template has no Page Performance page")

    slide_id = destination.slides._sldIdLst[-1]
    destination.slides._sldIdLst.remove(slide_id)
    destination.slides._sldIdLst.insert(page_performance_index + 1, slide_id)


def main() -> None:
    reference = Presentation(REFERENCE_DECK)
    reference_index = _slide_index(reference, "Traffic Acquisition")
    if reference_index is None:
        raise RuntimeError("Reference deck has no Traffic Acquisition page")
    reference_slide = reference.slides[reference_index]

    for template_path in sorted(TEMPLATE_DIR.glob("*.pptx")):
        presentation = Presentation(template_path)
        if _slide_index(presentation, "Traffic Acquisition") is not None:
            print(f"kept {template_path.name}: already present")
            continue
        _insert_reference_slide(presentation, reference_slide)
        presentation.save(template_path)
        print(f"updated {template_path.name}")


if __name__ == "__main__":
    main()
